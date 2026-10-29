"""Document parser service for the Legal Review App.

Extracts text from PDF, DOCX, and plain text files for legal review.
"""

from __future__ import annotations

import io
import logging
import os
from dataclasses import dataclass, field
from typing import Optional

import docx
import pdfplumber

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Parsed document data model
# ---------------------------------------------------------------------------

@dataclass
class ParsedDocument:
    """Holds extracted text, section structure, and metadata from a parsed document."""

    text: str
    sections: list[dict[str, str]] = field(default_factory=list)
    metadata: dict[str, object] = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Exceptions
# ---------------------------------------------------------------------------

class DocumentValidationError(Exception):
    """Raised when a file fails validation (format or size)."""


class DocumentParseError(Exception):
    """Raised when a document cannot be parsed or contains no extractable text."""


# ---------------------------------------------------------------------------
# Document parser
# ---------------------------------------------------------------------------

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


class DocumentParser:
    """Parses uploaded documents into structured text for legal review."""

    MAX_FILE_SIZE = MAX_FILE_SIZE

    # ------------------------------------------------------------------
    # Validation
    # ------------------------------------------------------------------

    def validate_file(self, filename: str, size: int) -> None:
        """Validate file extension and size.

        Raises ``DocumentValidationError`` if the file is not an accepted
        format or exceeds the maximum allowed size.
        """
        ext = os.path.splitext(filename)[1].lower()
        if ext not in ALLOWED_EXTENSIONS:
            accepted = ", ".join(sorted(ALLOWED_EXTENSIONS))
            raise DocumentValidationError(
                f"Unsupported file format '{ext}'. Accepted formats: {accepted}"
            )
        if size > self.MAX_FILE_SIZE:
            max_mb = self.MAX_FILE_SIZE / (1024 * 1024)
            raise DocumentValidationError(
                f"File size ({size} bytes) exceeds the maximum allowed size of {max_mb:.0f} MB."
            )

    # ------------------------------------------------------------------
    # Public parse entry-point
    # ------------------------------------------------------------------

    def parse(self, file_bytes: bytes, filename: str) -> ParsedDocument:
        """Parse *file_bytes* based on the file extension of *filename*.

        Returns a ``ParsedDocument`` with extracted text, sections, and
        metadata.  Raises ``DocumentParseError`` if the document contains
        no extractable text.
        """
        ext = os.path.splitext(filename)[1].lower()

        parsers = {
            ".pdf": self.parse_pdf,
            ".docx": self.parse_docx,
            ".txt": self.parse_text,
        }

        parser_fn = parsers.get(ext)
        if parser_fn is None:
            accepted = ", ".join(sorted(ALLOWED_EXTENSIONS))
            raise DocumentParseError(
                f"No parser available for extension '{ext}'. Accepted formats: {accepted}"
            )

        doc = parser_fn(file_bytes)

        if not doc.text or not doc.text.strip():
            raise DocumentParseError(
                "Document contains no extractable text."
            )

        # Attach common metadata
        doc.metadata["filename"] = filename
        doc.metadata["file_size"] = len(file_bytes)
        doc.metadata["content_type"] = ext.lstrip(".")

        return doc

    # ------------------------------------------------------------------
    # Format-specific parsers
    # ------------------------------------------------------------------

    def parse_pdf(self, file_bytes: bytes) -> ParsedDocument:
        """Extract text from a PDF using *pdfplumber*."""
        sections: list[dict[str, str]] = []
        pages_text: list[str] = []

        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page_num, page in enumerate(pdf.pages, start=1):
                    text = page.extract_text() or ""
                    pages_text.append(text)
                    if text.strip():
                        sections.append({
                            "title": f"Page {page_num}",
                            "content": text,
                        })
        except Exception as exc:
            raise DocumentParseError(f"Failed to parse PDF: {exc}") from exc

        full_text = "\n\n".join(pages_text)
        return ParsedDocument(
            text=full_text,
            sections=sections,
            metadata={"page_count": len(pages_text)},
        )

    def parse_docx(self, file_bytes: bytes) -> ParsedDocument:
        """Extract text from a DOCX file using *python-docx*, preserving
        section and paragraph structure."""
        sections: list[dict[str, str]] = []
        paragraphs_text: list[str] = []

        try:
            document = docx.Document(io.BytesIO(file_bytes))

            current_section_title: Optional[str] = None
            current_section_paragraphs: list[str] = []

            for para in document.paragraphs:
                text = para.text
                paragraphs_text.append(text)

                # Treat heading styles as section boundaries
                if para.style and para.style.name and para.style.name.startswith("Heading"):
                    # Save previous section if it has content
                    if current_section_title is not None and current_section_paragraphs:
                        sections.append({
                            "title": current_section_title,
                            "content": "\n".join(current_section_paragraphs),
                        })
                    current_section_title = text
                    current_section_paragraphs = []
                else:
                    current_section_paragraphs.append(text)

            # Flush last section
            if current_section_title is not None and current_section_paragraphs:
                sections.append({
                    "title": current_section_title,
                    "content": "\n".join(current_section_paragraphs),
                })
            elif not sections and current_section_paragraphs:
                # No headings found — treat entire document as one section
                sections.append({
                    "title": "Document",
                    "content": "\n".join(current_section_paragraphs),
                })

        except Exception as exc:
            raise DocumentParseError(f"Failed to parse DOCX: {exc}") from exc

        full_text = "\n".join(paragraphs_text)
        return ParsedDocument(
            text=full_text,
            sections=sections,
            metadata={"paragraph_count": len(paragraphs_text)},
        )

    def parse_text(self, file_bytes: bytes) -> ParsedDocument:
        """Read plain text content using UTF-8 encoding."""
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise DocumentParseError(
                f"Failed to decode text file as UTF-8: {exc}"
            ) from exc

        sections = [{"title": "Document", "content": text}] if text.strip() else []
        return ParsedDocument(
            text=text,
            sections=sections,
            metadata={},
        )
